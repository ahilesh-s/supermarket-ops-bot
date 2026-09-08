import json
from pathlib import Path
import sqlite3

import pytest

from abi_store import db, reports


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(db, 'ROOT', tmp_path)
    monkeypatch.setattr(db, 'DB_DIR', tmp_path / 'database')
    monkeypatch.setattr(db, 'STOCK_DB_PATH', tmp_path / 'database' / 'stock.db')
    monkeypatch.setattr(db, 'KHATA_DB_PATH', tmp_path / 'database' / 'khata.db')
    db.init_db()
    return tmp_path


def seed_bill(at, lines=None, mode='cash'):
    lines = lines or [dict(sku='A', name='Apple', qty=2, cost_price=30, line_total=100,
                           hsn_code='0808', unit='piece', mrp=50, taxable_value=100, cgst=0, sgst=0)]
    total = sum(line['line_total'] for line in lines)
    snapshot = dict(lines=lines, grand_total=total, total_tax=0, subtotal_taxable=total,
                    total_cgst=0, total_sgst=0, round_off=0, payable=total,
                    shop_name='Original & <Market>', shop_gstin=None)
    conn = db.get_connection()
    cur = conn.execute("INSERT INTO bills(status,created_at,finalized_at,payment_mode,totals_json) VALUES ('finalized',?,?,?,?)", (at,at,mode,json.dumps(snapshot)))
    conn.commit()
    conn.close()
    return cur.lastrowid


@pytest.mark.parametrize('start,end', [('2026-1-01','2026-01-02'), ('2026-02-30','2026-03-01'), ('2026-02-02','2026-02-01'), ('../bad','2026-01-01')])
def test_dates_are_strict(store, start, end):
    with pytest.raises(ValueError):
        reports.purchase_summary(start, end)


def test_business_day_utc_crossing(store):
    seed_bill('2026-01-01T18:29:59+00:00')
    seed_bill('2026-01-01T18:30:00+00:00')
    seed_bill('2026-01-02T18:30:00+00:00')
    seed_bill('2026-01-02T20:00:00+00:00')
    assert reports.daily_close('2026-01-02')['bill_count'] == 1
    from abi_store.preferences import set_preference
    set_preference('timezone', 'America/New_York')
    assert reports.daily_close('2026-01-01')['bill_count'] == 2


def test_period_ranking_uses_all_skus(store):
    for day in (1, 2):
        lines = [dict(sku=f'{day}-{i}', name='Same name', qty=1, line_total=100, cost_price=20) for i in range(5)]
        lines.append(dict(sku='consistent', name='Always sixth', qty=1, line_total=90, cost_price=10))
        seed_bill(f'2026-01-0{day}T10:00:00+00:00', lines)
    result = reports.sales_summary('2026-01-01', '2026-01-02')
    assert result['top_items'][0] == dict(sku='consistent', name='Always sixth', qty=2, revenue=180)
    assert len(result['items']) == 11


def test_margin_uses_snapshot_cost_and_refuses_legacy(store):
    bill_id = seed_bill('2026-01-01T10:00:00+00:00')
    result = reports.margin_report('2026-01-01', '2026-01-01')
    assert (result['revenue'], result['cost'], result['margin']) == (100, 60, 40)
    assert 'not accounting net profit' in result['basis']
    conn = db.get_connection()
    snapshot = json.loads(conn.execute('SELECT totals_json FROM bills WHERE id=?', (bill_id,)).fetchone()[0])
    del snapshot['lines'][0]['cost_price']
    conn.execute('UPDATE bills SET totals_json=? WHERE id=?', (json.dumps(snapshot),bill_id))
    conn.commit()
    conn.close()
    with pytest.raises(ValueError, match='cost'):
        reports.margin_report('2026-01-01', '2026-01-01')


def test_paid_modes_only_and_repayments_separate(store):
    for mode in ('cash', 'upi', 'card', 'bank_transfer', 'credit'):
        seed_bill('2026-01-01T10:00:00+00:00', mode=mode)
    conn = db.get_khata_connection()
    conn.execute("INSERT INTO khata_txns(customer,amount,txn_type,created_at) VALUES ('Test',1000,'payment','2026-01-01T10:00:00+00:00')")
    conn.commit()
    conn.close()
    result = reports.daily_close('2026-01-01')
    assert result['total_sales'] == 400
    assert result['bill_count'] == 4
    assert result['excluded_bill_count'] == 1


def test_invoice_snapshot_escaping_a4_and_runtime_root(store):
    from abi_store import documents
    from abi_store.preferences import set_preference
    from pypdf import PdfReader
    bill_id = seed_bill('2026-01-01T20:00:00+00:00')
    conn = db.get_connection()
    data = json.loads(conn.execute('SELECT totals_json FROM bills WHERE id=?', (bill_id,)).fetchone()[0])
    data['lines'][0]['name'] = 'Long & <special> ' + 'Wrapped product description ' * 15
    conn.execute('UPDATE bills SET totals_json=? WHERE id=?', (json.dumps(data),bill_id))
    conn.commit()
    conn.close()
    set_preference('shop_name', 'Changed shop')
    set_preference('shop_gstin', 'CHANGED')
    path = Path(documents.generate_invoice_pdf(bill_id))
    assert path.is_relative_to(store / 'generated')
    reader = PdfReader(path)
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert 'Original & <Market>' in text
    assert 'Changed shop' not in text and 'CHANGED' not in text
    assert 'not a compliant tax invoice' in text.lower()
    assert 'Long & <special>' in text
    assert '02 Jan 2026' in text
    assert len(reader.pages) == 1
    assert float(reader.pages[0].mediabox.width) == pytest.approx(595.28, abs=.1)


def test_documents_import_has_no_filesystem_side_effects(store, monkeypatch):
    import importlib
    import os
    from abi_store import documents
    def forbidden(*args, **kwargs):
        pytest.fail('Import attempted directory creation')
    monkeypatch.setattr(Path, 'mkdir', forbidden)
    monkeypatch.setattr(os, 'makedirs', forbidden)
    importlib.reload(documents)
    assert not (store / 'generated').exists()


def test_deck_real_data_unique_outputs_and_empty_period(store):
    from abi_store import documents
    from pptx import Presentation
    from zipfile import ZipFile
    seed_bill('2026-01-01T10:00:00+00:00')
    first = Path(documents.generate_analysis_deck('2026-01-01', '2026-01-02'))
    second = Path(documents.generate_analysis_deck('2026-01-01', '2026-01-02'))
    assert first != second and first.is_relative_to(store / 'generated')
    prs = Presentation(first)
    text = '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    assert '100.00' in text and 'Current stock' in text and 'not historical' in text
    assert 'Khata repayments excluded' in text
    charts = [shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]
    assert charts and list(charts[0].series[0].values) == [100, 0]
    with ZipFile(first) as archive:
        assert archive.testzip() is None
        assert any(name.startswith('ppt/charts/') for name in archive.namelist())
    empty = Presentation(documents.generate_analysis_deck('2025-01-01', '2025-01-01'))
    assert 'No finalized paid sales' in '\n'.join(s.text for sl in empty.slides for s in sl.shapes if s.has_text_frame)


def test_deck_rejects_dates_before_creating_paths(store):
    from abi_store import documents
    with pytest.raises(ValueError):
        documents.generate_analysis_deck('../bad', '2026-01-01')
    assert not (store / 'generated').exists()


def test_microsecond_boundaries_are_exact(store):
    seed_bill('2026-01-01T18:29:59.999999+00:00')
    seed_bill('2026-01-01T18:30:00.000000+00:00')
    seed_bill('2026-01-02T18:29:59.999999+00:00')
    result = reports.sales_summary('2026-01-02', '2026-01-02')
    assert result['bill_count'] == 2
    assert result['daily_totals'][0]['bill_count'] == 2


def test_purchase_and_sales_follow_dst_business_day(store):
    from abi_store.preferences import set_preference
    set_preference('timezone', 'America/New_York')
    conn = db.get_connection()
    conn.execute("INSERT INTO products(sku,name,unit,hsn_code,gst_rate,cost_price,mrp) VALUES ('A','Apple','piece','0808',0,999,1000)")
    stamps = ['2026-03-08T04:59:59+00:00', '2026-03-08T05:00:00Z',
              '2026-03-09T03:59:59+00:00', '2026-03-09T04:00:00+00:00']
    for at in stamps:
        conn.execute("INSERT INTO stock_receipts(sku,qty,cost_price,mrp,received_at) VALUES ('A',1,10,50,?)", (at,))
    conn.commit()
    conn.close()
    for at in stamps:
        seed_bill(at)
    assert reports.purchase_summary('2026-03-08','2026-03-08')['total_cost'] == 20
    assert reports.daily_close('2026-03-08')['bill_count'] == 2
    assert reports.margin_report('2026-03-08','2026-03-08')['cost'] == 120


def test_invalid_timezone_fails_instead_of_falling_back(store):
    from abi_store.preferences import set_preference
    set_preference('timezone', 'Not/AZone')
    with pytest.raises(ValueError, match='timezone'):
        reports.daily_close('2026-01-01')


def test_period_ranking_reaches_deck_chart(store):
    from abi_store.documents import generate_analysis_deck
    from pptx import Presentation
    for day in (1, 2):
        lines = [dict(sku=f'{day}-{i}', name=f'Competitor {day}-{i}', qty=1, line_total=100, cost_price=20) for i in range(5)]
        lines.append(dict(sku='consistent', name='Always sixth', qty=1, line_total=90, cost_price=10))
        seed_bill(f'2026-01-0{day}T10:00:00+00:00', lines)
    deck = Presentation(generate_analysis_deck('2026-01-01','2026-01-02'))
    charts = [sh.chart for sl in deck.slides for sh in sl.shapes if sh.has_chart]
    assert charts[1].plots[0].categories[0].label == 'Always sixth [consistent]'
    assert charts[1].series[0].values[0] == 180


def test_long_period_charts_preserve_all_sales(store):
    from abi_store.documents import generate_analysis_deck
    from pptx import Presentation
    seed_bill('2026-01-01T10:00:00+00:00')
    seed_bill('2026-12-31T10:00:00+00:00')
    deck = Presentation(generate_analysis_deck('2026-01-01','2026-12-31'))
    chart = next(sh.chart for sl in deck.slides for sh in sl.shapes if sh.has_chart)
    assert len(chart.series[0].values) <= 31
    assert sum(chart.series[0].values) == 200
