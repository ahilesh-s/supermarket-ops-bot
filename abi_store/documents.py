"""Snapshot invoices and fresh, data-backed presentations; no import-time I/O."""
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape
import textwrap

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE

from . import db
from .billing import get_finalized_bill
from .reports import sales_summary, business_timezone, validate_date_range
from .inventory import low_stock_items


def _rs(amount: float) -> str:
    return f"Rs. {amount:,.2f}"


def _output_path(kind: str, stem: str, extension: str) -> Path:
    directory = db.ROOT / "generated" / kind
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{stem}_{uuid4().hex}.{extension}"


def generate_invoice_pdf(bill_id: int) -> str:
    """Render finalized prices and shop identity, not today's preferences."""
    if isinstance(bill_id, bool) or not isinstance(bill_id, int) or bill_id <= 0:
        raise ValueError("bill_id must be a positive integer")
    bill = get_finalized_bill(bill_id)
    at = datetime.fromisoformat(bill["finalized_at"])
    if at.tzinfo is None:
        at = at.replace(tzinfo=timezone.utc)
    zone = business_timezone()
    stamp = at.astimezone(zone).strftime("%d %b %Y, %I:%M %p")
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("InvoiceCell", parent=styles["Normal"], fontSize=7, leading=9, splitLongWords=True)
    header = ParagraphStyle("InvoiceHeader", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")

    def para(value, style=None):
        return Paragraph(escape(str(value)), style or styles["Normal"])

    gstin = bill.get("shop_gstin")
    has_gstin = bool(gstin and str(gstin).strip() not in {"-", "None"})
    story = [para(bill.get("shop_name") or "Shop identity unavailable in legacy snapshot", styles["Title"]),
             para("Sales invoice", styles["Heading2"]),
             para(f"GSTIN: {gstin}" if has_gstin else "GSTIN unavailable: not a compliant tax invoice."),
             para("Generated sales record; GST compliance is not certified by this software."),
             para(f"Invoice #{bill['bill_id']} | {stamp} ({zone})"),
             para(f"Payment: {bill['payment_mode']}"), Spacer(1, 6 * mm)]
    rows = [[para(h, header) for h in ("Item", "HSN", "Qty", "Rate", "Taxable", "CGST", "SGST", "Total")]]
    for line in bill["lines"]:
        values = [line["name"], line["hsn_code"], f"{line['qty']:g} {line['unit']}",
                  _rs(line["mrp"]), _rs(line["taxable_value"]), _rs(line["cgst"]),
                  _rs(line["sgst"]), _rs(line["line_total"])]
        rows.append([para(value, cell) for value in values])
    # 180 mm fits A4 with explicit 15 mm side margins. Paragraphs wrap every field.
    table = Table(rows, colWidths=[v * mm for v in (44, 16, 19, 20, 22, 18, 18, 23)], repeatRows=1,
                  splitInRow=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173342")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6f7")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, 0), (-1, 0), .5, colors.grey),
    ]))
    story.extend([table, Spacer(1, 6 * mm)])
    totals = [("Subtotal (taxable value)", "subtotal_taxable"), ("Total CGST", "total_cgst"),
              ("Total SGST", "total_sgst"), ("Round off", "round_off"), ("Payable", "payable")]
    summary = Table([[para(label), para(_rs(bill[key]))] for label, key in totals], colWidths=[130 * mm, 50 * mm])
    summary.setStyle(TableStyle([("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#173342"))]))
    story.append(summary)
    path = _output_path("invoices", f"invoice_{bill_id}", "pdf")
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm, title=f"Sales invoice {bill_id}")
    doc.build(story)
    return str(path)


INK = RGBColor(23, 51, 66)
TEAL = RGBColor(0, 137, 123)
MUTED = RGBColor(88, 106, 118)


def _textbox(slide, text, x, y, w, h, size=18, bold=False, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = text
    for paragraph in frame.paragraphs:
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
    return box


def _chart(slide, categories, values, title, horizontal=False):
    # Native charts embed a private workbook in each fresh presentation: no shared PNG paths.
    data = CategoryChartData()
    data.categories = categories
    data.add_series("Tax-inclusive sales (Rs.)", values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(.7), Inches(1.45), Inches(11.9), Inches(4.8), data).chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(16)
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.value_axis.tick_labels.font.size = Pt(11)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.tick_labels.number_format = '#,##0'
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = TEAL


def generate_analysis_deck(start_date: str, end_date: str) -> str:
    """Create a new 16:9 deck from all finalized paid-sale SKU snapshots."""
    validate_date_range(start_date, end_date)  # must precede paths and queries
    data = sales_summary(start_date, end_date)
    stock = low_stock_items()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    period = f"{start_date} to {end_date} | {data['timezone']}"

    def slide(title, subtitle=period):
        page = prs.slides.add_slide(prs.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor(247, 249, 250)
        _textbox(page, title, .6, .35, 12.1, .6, 28, True)
        _textbox(page, subtitle, .6, .99, 12.1, .35, 11, color=MUTED)
        _textbox(page, f"Supermarket Ops | {len(prs.slides)}", .6, 7.08, 12, .22, 9, color=MUTED)
        return page

    overview = slide("Sales analysis")
    for i, (label, value) in enumerate((("Tax-inclusive sales", _rs(data["total_sales"])),
                                       ("Recorded tax component", _rs(data["total_tax"])),
                                       ("Finalized paid bills", str(data["bill_count"])))):
        x = .65 + i * 4.2
        panel = overview.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(3.95), Inches(1.5))
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
        panel.line.color.rgb = RGBColor(218, 228, 232)
        _textbox(overview, label, x + .15, 1.8, 3.65, .4, 14, color=MUTED)
        _textbox(overview, value, x + .15, 2.3, 3.65, .6, 25, True)
    modes = " | ".join(f"{mode}: {_rs(amount)}" for mode, amount in data['by_payment_mode'].items())
    _textbox(overview, modes or "No finalized paid sales", .7, 3.55, 12, 1.1, 18)
    _textbox(overview, "Khata repayments excluded. Credit sales are unsupported.\n"
             f"Excluded unsupported-mode bills: {data['excluded_bill_count']}\n"
             "Sales use recorded grand totals before payment rounding, not bank settlement totals.",
             .7, 5, 12, 1.45, 16, color=MUTED)

    trend = slide("Sales trend")
    if data['bill_count']:
        days = data['daily_totals']
        # Bound chart density: sum contiguous buckets, never sample away revenue.
        size = max(1, (len(days) + 30) // 31)
        groups = [days[i:i + size] for i in range(0, len(days), size)]
        labels = [g[0]['date'] if len(g) == 1 else f"{g[0]['date']} to {g[-1]['date']}" for g in groups]
        _chart(trend, labels, [sum(d['total_sales'] for d in g) for g in groups],
               "Daily sales (Rs.)" if size == 1 else f"Sales in up to {size}-day buckets (Rs.)")
    else:
        _textbox(trend, "No finalized paid sales", .8, 2.6, 11, 1, 26)

    top = slide("Top items by full-period revenue", period + " | Ranked by SKU, not daily top-five lists")
    if data['top_items']:
        labels = [textwrap.shorten(f"{i['name']} [{i['sku']}]", width=55, placeholder="...") for i in data['top_items']]
        _chart(top, labels, [i['revenue'] for i in data['top_items']], "Top five SKUs (Rs.)", horizontal=True)
    else:
        _textbox(top, "No finalized paid sales", .8, 2.6, 11, 1, 26)

    generated = datetime.now(timezone.utc).isoformat(timespec='seconds')
    chunks = [stock[i:i + 8] for i in range(0, len(stock), 8)] or [[]]
    for chunk in chunks:
        current = slide("Current stock | reorder review", f"Current at generation ({generated}); not historical period-end stock")
        if not chunk:
            _textbox(current, "No items below reorder level.", .8, 2, 11.7, 1, 22)
        for i, item in enumerate(chunk):
            name = textwrap.shorten(str(item['name']), width=85, placeholder="...")
            _textbox(current, f"{name}: {item['qty']:g} {item.get('unit', '')} left | reorder at {item['reorder_level']:g}",
                     .8, 1.65 + i * .6, 11.7, .58, 17)
    path = _output_path("presentations", f"analysis_{start_date}_to_{end_date}", "pptx")
    prs.save(str(path))
    return str(path)
